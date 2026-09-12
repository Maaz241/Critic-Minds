"""
pptx.py
Extracts text from PowerPoint presentations (PPTX), preserving slide numbers,
slide titles, text frames, bullet points, tables, and speaker notes for citation.
"""
import io
from typing import List, Tuple
from pptx import Presentation


class PPTXExtractionError(Exception):
    pass


def extract_pptx_slides(file_bytes: bytes) -> List[Tuple[int, str, str]]:
    """
    Extracts slides from PPTX bytes.
    Returns a list of tuples: (slide_number, slide_title, slide_content).
    Slide numbering is 1-indexed.
    """
    try:
        prs = Presentation(io.BytesIO(file_bytes))
    except Exception as exc:
        raise PPTXExtractionError(f"Failed to open PPTX presentation: {exc}") from exc

    slides: List[Tuple[int, str, str]] = []

    for i, slide in enumerate(prs.slides):
        slide_num = i + 1
        title = f"Slide {slide_num}"
        content_items: List[str] = []

        # Find shapes with text
        for shape in slide.shapes:
            if shape.has_text_frame:
                # If shape is title
                if shape == slide.shapes.title and shape.text.strip():
                    title = shape.text.strip()
                else:
                    text = shape.text.strip()
                    if text:
                        content_items.append(text)
            elif shape.has_table:
                table_cells = []
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        table_cells.append(" | ".join(cells))
                if table_cells:
                    content_items.append("\n".join(table_cells))

        # Check speaker notes if present
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                content_items.append(f"[Speaker Notes]: {notes}")

        combined_text = "\n\n".join(content_items)
        if combined_text or (title and title != f"Slide {slide_num}"):
            slides.append((slide_num, title, f"{title}\n{combined_text}".strip()))

    if not slides:
        raise PPTXExtractionError("No readable slides or text content found in this PPTX presentation.")

    return slides
