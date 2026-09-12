"""
docx.py
Extracts text from DOCX documents, tracking headings, sections, and table contents.
"""
import io
from typing import List, Tuple
from docx import Document


class DOCXExtractionError(Exception):
    pass


def extract_docx_sections(file_bytes: bytes) -> List[Tuple[str, str]]:
    """
    Extracts text sections from a DOCX file.
    Returns a list of tuples: (section_title, text_content).
    """
    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise DOCXExtractionError(f"Failed to open DOCX document: {exc}") from exc

    sections: List[Tuple[str, str]] = []
    current_heading = "General"
    current_paragraphs: List[str] = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue

        if p.style and p.style.name and p.style.name.startswith("Heading"):
            if current_paragraphs:
                sections.append((current_heading, "\n\n".join(current_paragraphs)))
                current_paragraphs = []
            current_heading = text
        else:
            current_paragraphs.append(text)

    # Process tables
    table_texts: List[str] = []
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_texts.append(" | ".join(cells))

    if table_texts:
        current_paragraphs.append("\n".join(table_texts))

    if current_paragraphs:
        sections.append((current_heading, "\n\n".join(current_paragraphs)))

    if not sections:
        raise DOCXExtractionError("This DOCX document appears to be empty.")

    return sections
